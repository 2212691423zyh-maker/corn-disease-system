from __future__ import annotations

import hashlib
import os
import uuid
from typing import Optional, List, Dict, Any, Tuple

import fitz

from ..config import settings
from ..rag.retriever import HybridRetriever
from ..database import SessionLocal
from ..models.knowledge import KnowledgeDocument, KnowledgeChunk


class KnowledgeService:
    """知识库文档处理服务 —— 上传 → 解析 → 分块 → 向量化 → 持久化"""

    CHUNK_SIZE = 500
    CHUNK_OVERLAP = 80

    @classmethod
    def save_upload(cls, filename: str, content: bytes) -> Tuple[str, str, int]:
        """保存上传文件，返回 (file_path, doc_type, file_size)"""
        ext = os.path.splitext(filename)[1].lower()
        ext_clean = ext.lstrip(".")
        doc_type = ext_clean if ext_clean in {"pdf", "doc", "docx", "txt"} else "txt"
        new_name = f"{uuid.uuid4().hex}{ext}"
        file_path = os.path.join(settings.UPLOAD_DIR, "documents", new_name)
        os.makedirs(os.path.dirname(file_path), exist_ok=True)
        with open(file_path, "wb") as f:
            f.write(content)
        return file_path, doc_type, len(content)

    @classmethod
    def extract_text(cls, file_path: str, doc_type: str) -> str:
        """从文件中提取纯文本"""
        if doc_type == "pdf":
            return cls._extract_pdf(file_path)
        elif doc_type in {"doc", "docx"}:
            return cls._extract_docx(file_path)
        else:
            try:
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    return f.read()
            except Exception as e:
                raise ValueError(f"文本文件读取失败: {e}") from e

    @staticmethod
    def _extract_pdf(path: str) -> str:
        try:
            with open(path, "rb") as f:
                magic = f.read(8)
        except OSError as e:
            raise ValueError(f"无法读取文件: {e}") from e

        if not magic.startswith(b"%PDF"):
            raise ValueError("PDF文件结构损坏或不是有效的PDF格式，无法解析")

        try:
            doc = fitz.open(path)
        except Exception as e:
            raise ValueError(f"PDF文件结构损坏，无法解析: {e}") from e

        try:
            text_parts = []
            for page in doc:
                try:
                    text_parts.append(page.get_text())
                except Exception as e:
                    text_parts.append("")
            text = "\n".join(text_parts)
        finally:
            doc.close()

        if not text.strip():
            raise ValueError("PDF文件内容为空或无法提取文本（可能是扫描件/图片型PDF）")
        return text

    @staticmethod
    def _extract_docx(path: str) -> str:
        try:
            from docx import Document
            from docx.opc.exceptions import PackageNotFoundError
        except ImportError as e:
            raise ValueError("缺少 python-docx 库，无法解析 Word 文档") from e

        try:
            doc = Document(path)
        except PackageNotFoundError:
            raise ValueError("Word文件结构损坏或不是有效的docx格式，无法解析")
        except Exception as e:
            raise ValueError(f"Word文件解析失败: {e}") from e

        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())

    @staticmethod
    def _cleanup_file(file_path: str) -> None:
        try:
            if os.path.exists(file_path):
                os.remove(file_path)
        except Exception:
            pass

    @classmethod
    def chunk_text(cls, text: str, doc_name: str, document_id: int) -> List[Dict[str, Any]]:
        """按 CHUNK_SIZE + OVERLAP 分块"""
        chunks = []
        start = 0
        text_len = len(text)
        idx = 0
        while start < text_len:
            end = min(start + cls.CHUNK_SIZE, text_len)
            chunk_text = text[start:end].strip()
            if chunk_text:
                chunk_id = hashlib.md5(f"{doc_name}:{idx}:{chunk_text[:50]}".encode()).hexdigest()
                chunks.append({
                    "id": chunk_id,
                    "text": chunk_text,
                    "document_id": document_id,
                    "doc_name": doc_name,
                    "chunk_index": idx,
                    "token_count": len(chunk_text),
                })
            idx += 1
            start += cls.CHUNK_SIZE - cls.CHUNK_OVERLAP
            if end >= text_len:
                break
        return chunks

    @classmethod
    def index_document(cls, filename: str, content: bytes, document_id: int) -> dict:
        """完整处理流程：保存 → 解析 → 分块 → 向量化 → 索引 → 持久化"""
        file_path, doc_type, file_size = cls.save_upload(filename, content)
        try:
            text = cls.extract_text(file_path, doc_type)
        except ValueError as e:
            cls._cleanup_file(file_path)
            return {"status": "failed", "error": str(e)}
        except Exception as e:
            cls._cleanup_file(file_path)
            return {"status": "failed", "error": f"文档解析失败: {e}"}

        if not text.strip():
            return {"status": "failed", "error": "文档内容为空或无法提取文本"}

        chunks = cls.chunk_text(text, filename, document_id)

        retriever = HybridRetriever.get()
        retriever.add_chunks(chunks)

        cls._persist_to_db(filename, file_path, doc_type, file_size, chunks)

        return {
            "status": "indexed",
            "chunk_count": len(chunks),
            "file_size": file_size,
            "doc_type": doc_type,
            "document_id": document_id,
        }

    @staticmethod
    def _persist_to_db(filename: str, file_path: str, doc_type: str, file_size: int, chunks: List[Dict[str, Any]]):
        try:
            db = SessionLocal()
            doc = KnowledgeDocument(
                filename=filename,
                file_path=file_path,
                doc_type=doc_type,
                chunk_count=len(chunks),
                file_size=file_size,
                status="indexed",
            )
            db.add(doc)
            db.flush()
            for c in chunks:
                db.add(KnowledgeChunk(
                    document_id=doc.id,
                    chunk_index=c["chunk_index"],
                    chunk_id=c["id"],
                    content=c["text"],
                    token_count=c.get("token_count", 0),
                ))
            db.commit()
            db.close()
        except Exception as e:
            print(f"[KnowledgeService] persist_to_db error: {e}")

    @classmethod
    def list_documents(cls) -> List[Dict[str, Any]]:
        db = SessionLocal()
        try:
            docs = db.query(KnowledgeDocument).order_by(KnowledgeDocument.created_at.desc()).all()
            return [
                {
                    "id": d.id,
                    "filename": d.filename,
                    "doc_type": d.doc_type,
                    "chunk_count": d.chunk_count,
                    "file_size": d.file_size,
                    "status": d.status,
                    "created_at": d.created_at.isoformat() if d.created_at else None,
                }
                for d in docs
            ]
        finally:
            db.close()

    @classmethod
    def delete_document(cls, document_id: int) -> bool:
        db = SessionLocal()
        try:
            doc = db.query(KnowledgeDocument).filter(KnowledgeDocument.id == document_id).first()
            if not doc:
                return False
            if os.path.exists(doc.file_path):
                try:
                    os.remove(doc.file_path)
                except Exception:
                    pass
            db.delete(doc)
            db.commit()
            return True
        except Exception:
            db.rollback()
            return False
        finally:
            db.close()
