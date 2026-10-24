from __future__ import annotations

import io
import os
import uuid
from typing import Optional

from fastapi import HTTPException

from ..ai.client import create_llm_client
from ..rag.retriever import HybridRetriever
from ..config import settings


PLAN_SYSTEM_PROMPT = """你是玉米病害防治专家，请基于给定的识别结果和 RAG 检索证据，生成一份结构化的防治方案。

必须严格返回以下 JSON 格式（不要任何多余文字）：
{
    "plan_title": "方案标题",
    "medication_plan": [
        {"drug": "药剂名称", "dosage": "使用浓度/剂量", "frequency": "使用频率", "precautions": "注意事项"}
    ],
    "fertilization_plan": [
        {"type": "肥料类型", "dosage": "用量", "timing": "施用时期"}
    ],
    "prevention_tips": "综合防治建议（Markdown 格式，分点）",
    "agricultural_operations": "农业操作建议（深翻/密植/排水等）"
}
"""


class PlanService:
    """防治方案生成与导出"""

    @classmethod
    def generate(cls, disease_name: str, symptoms: str = "", rag_evidence: Optional[list] = None) -> dict:
        # ---- 1. RAG 检索 ----
        retriever = HybridRetriever.get()
        rag_chunks = retriever.search(
            f"{disease_name} 防治方法 药剂 农业措施",
            top_k=5,
        )

        evidence_text = "\n".join(
            f"来源: {c.get('source', '')}\n{c['text']}" for c in rag_chunks
        ) or "(未检索到相关知识)"

        # ---- 2. LLM 生成方案 ----
        llm = create_llm_client()
        messages = [
            {"role": "system", "content": PLAN_SYSTEM_PROMPT},
            {
                "role": "user",
                "content": f"病害名称: {disease_name}\n症状描述: {symptoms}\n\n【RAG 检索证据】\n{evidence_text}",
            },
        ]

        try:
            plan_json = llm.chat_json(messages, temperature=0.3)
        except Exception as e:
            raise HTTPException(502, f"方案生成失败: {e}")

        # ---- 3. 附加 RAG 证据 ----
        plan_json["rag_evidence"] = [
            {
                "source": c.get("source", ""),
                "citation": c["text"][:200],
                "score": round(c.get("score", 0), 3),
            }
            for c in rag_chunks
        ]
        plan_json["model_used"] = llm.name

        return plan_json

    @classmethod
    def export_word(cls, plan: dict, plan_id: int) -> dict:
        """导出 Word 文档（后端 python-docx）"""
        from docx import Document
        from docx.shared import Pt

        doc = Document()

        title = plan.get("plan_title", "玉米病害防治方案")
        doc.add_heading(title, level=0)

        doc.add_heading("一、病情诊断", level=1)
        primary = plan.get("primary", {})
        doc.add_paragraph(f"病害名称：{primary.get('name', '未知')}")
        doc.add_paragraph(f"置信度：{primary.get('confidence', 0):.1%}")
        if primary.get("symptoms"):
            doc.add_paragraph(f"症状描述：{primary['symptoms']}")

        doc.add_heading("二、用药方案", level=1)
        meds = plan.get("medication_plan", [])
        if meds:
            table = doc.add_table(rows=1, cols=4)
            table.style = "Light Grid Accent 1"
            hdr = table.rows[0].cells
            hdr[0].text = "药剂"
            hdr[1].text = "剂量"
            hdr[2].text = "频率"
            hdr[3].text = "注意事项"
            for m in meds:
                row = table.add_row().cells
                row[0].text = m.get("drug", "")
                row[1].text = m.get("dosage", "")
                row[2].text = m.get("frequency", "")
                row[3].text = m.get("precautions", "")
        else:
            doc.add_paragraph("暂无用药建议")

        doc.add_heading("三、施肥方案", level=1)
        for f in plan.get("fertilization_plan", []):
            doc.add_paragraph(f"• {f.get('type', '')} - {f.get('dosage', '')} - {f.get('timing', '')}")

        doc.add_heading("四、综合防治建议", level=1)
        if plan.get("prevention_tips"):
            doc.add_paragraph(plan["prevention_tips"])
        if plan.get("agricultural_operations"):
            doc.add_paragraph(plan["agricultural_operations"])

        # ---- 保存 ----
        exports_dir = os.path.join(settings.UPLOAD_DIR, "exports")
        os.makedirs(exports_dir, exist_ok=True)
        filename = f"plan_{plan_id}_{uuid.uuid4().hex[:8]}.docx"
        file_path = os.path.join(exports_dir, filename)
        doc.save(file_path)

        file_size = os.path.getsize(file_path)
        return {"file_path": file_path, "filename": filename, "file_size": file_size}
